"""Build the Russian user guide DOCX from the canonical Markdown source."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "USER_GUIDE_RU.md"
OUTPUT = ROOT / "Руководство_пользователя_RAG_Assistant.docx"

NAVY = "17324D"
BLUE = "176B87"
PALE_BLUE = "EAF4F7"
PALE_GRAY = "F3F5F7"
MID_GRAY = "66717D"
WHITE = "FFFFFF"


def shade(cell, color: str) -> None:
    props = cell._tc.get_or_add_tcPr()
    element = props.find(qn("w:shd"))
    if element is None:
        element = OxmlElement("w:shd")
        props.append(element)
    element.set(qn("w:fill"), color)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120) -> None:
    props = cell._tc.get_or_add_tcPr()
    tc_mar = props.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        props.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    props = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    props.append(marker)


def prevent_row_split(row) -> None:
    props = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:cantSplit")
    marker.set(qn("w:val"), "true")
    props.append(marker)


def set_run_font(run, name="Aptos", size=None, bold=None, color=None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def add_hyperlink(paragraph, label: str, url: str) -> None:
    rel_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.extend((color, underline))
    text = OxmlElement("w:t")
    text.text = label
    run.extend((props, text))
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


INLINE = re.compile(r"(\[[^\]]+\]\([^\)]+\)|`[^`]+`|\*\*[^*]+\*\*)")


def add_inline(paragraph, text: str, *, default_bold=False, default_color=None) -> None:
    cursor = 0
    for match in INLINE.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor : match.start()])
            set_run_font(run, bold=default_bold, color=default_color)
        token = match.group(0)
        if token.startswith("["):
            label, url = token[1:].split("](", 1)
            add_hyperlink(paragraph, label, url[:-1])
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, name="Cascadia Mono", size=9.2, color=NAVY)
            run._element.get_or_add_rPr().append(OxmlElement("w:shd"))
            run._element.rPr[-1].set(qn("w:fill"), PALE_GRAY)
        else:
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, bold=True, color=default_color)
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        set_run_font(run, bold=default_bold, color=default_color)


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.4)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    normal.paragraph_format.line_spacing = 1.08

    for style_name, size, before, after in (
        ("Heading 1", 18, 18, 7),
        ("Heading 2", 13, 13, 5),
        ("Heading 3", 11, 10, 4),
    ):
        style = doc.styles[style_name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(NAVY if style_name == "Heading 1" else BLUE)
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = doc.styles[name]
        style.font.name = "Aptos"
        style.font.size = Pt(10.2)
        style.paragraph_format.space_after = Pt(3)
        style.paragraph_format.left_indent = Cm(0.62)
        style.paragraph_format.first_line_indent = Cm(-0.35)


def add_title_page(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.1)
    section.right_margin = Cm(2.1)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(75)
    p.paragraph_format.space_after = Pt(14)
    run = p.add_run("RAG ASSISTANT")
    set_run_font(run, name="Aptos Display", size=12, bold=True, color=BLUE)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(12)
    run = p.add_run("Подробная инструкция\nпользователя")
    set_run_font(run, name="Aptos Display", size=30, bold=True, color=NAVY)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(30)
    run = p.add_run(
        "Проекты и чаты, общие и приватные документы, поиск по источникам, "
        "очереди, формулы и безопасная работа"
    )
    set_run_font(run, size=13, color=MID_GRAY)

    table = doc.add_table(rows=3, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    table.columns[0].width = Cm(13.7)
    for cell in table.column_cells(0):
        shade(cell, PALE_BLUE)
        set_cell_margins(cell, 130, 180, 130, 180)
    for cell, label, value in zip(
        table.column_cells(0),
        ("СРЕДА", "МОДЕЛЬ ДОСТУПА", "АКТУАЛЬНОСТЬ"),
        (
            "https://5.35.120.112:8443/",
            "Персональная учётная запись; общие документы и приватные проекты",
            "Редакция после внедрения локальной регистрации и авторизации",
        ),
    ):
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(label + "\n")
        set_run_font(r, size=8, bold=True, color=BLUE)
        r = p.add_run(value)
        set_run_font(r, size=10.2, color=NAVY)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(45)
    r = p.add_run("НЭК · 2026")
    set_run_font(r, size=9.5, bold=True, color=MID_GRAY)
    doc.add_page_break()


def add_footer(doc: Document) -> None:
    for section in doc.sections:
        footer = section.footer
        p = footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = p.add_run("RAG Assistant · инструкция пользователя   ")
        set_run_font(r, size=8.3, color=MID_GRAY)
        fld = OxmlElement("w:fldSimple")
        fld.set(qn("w:instr"), "PAGE")
        p._p.append(fld)


def parse_table(lines: list[str], start: int, doc: Document) -> int:
    rows = []
    i = start
    while i < len(lines) and lines[i].strip().startswith("|"):
        rows.append([part.strip() for part in lines[i].strip().strip("|").split("|")])
        i += 1
    if len(rows) < 2:
        return start
    headers = rows[0]
    body = rows[2:]
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_repeat_table_header(table.rows[0])
    prevent_row_split(table.rows[0])
    for cell, value in zip(table.rows[0].cells, headers):
        shade(cell, NAVY)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        add_inline(cell.paragraphs[0], value, default_bold=True, default_color=WHITE)
    for row_index, values in enumerate(body):
        row = table.add_row()
        prevent_row_split(row)
        cells = row.cells
        for cell, value in zip(cells, values):
            if row_index % 2:
                shade(cell, PALE_GRAY)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            add_inline(cell.paragraphs[0], value)
            cell.paragraphs[0].paragraph_format.space_after = Pt(0)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return i


def add_code_block(doc: Document, content: list[str]) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    cell = table.cell(0, 0)
    shade(cell, PALE_GRAY)
    set_cell_margins(cell, 120, 150, 120, 150)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    for idx, line in enumerate(content):
        if idx:
            p.add_run().add_break()
        r = p.add_run(line)
        set_run_font(r, name="Cascadia Mono", size=8.7, color=NAVY)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_blockquote(doc: Document, text: str) -> None:
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    table.columns[0].width = Cm(0.12)
    table.columns[1].width = Cm(15.0)
    shade(table.cell(0, 0), BLUE)
    shade(table.cell(0, 1), PALE_BLUE)
    set_cell_margins(table.cell(0, 1), 110, 160, 110, 160)
    p = table.cell(0, 1).paragraphs[0]
    add_inline(p, text)
    p.paragraph_format.space_after = Pt(0)


def build() -> None:
    text = SOURCE.read_text(encoding="utf-8")
    lines = text.splitlines()
    doc = Document()
    configure_styles(doc)
    add_title_page(doc)

    index = 1  # Markdown title is represented by the designed title page.
    paragraph_buffer: list[str] = []

    def flush_buffer() -> None:
        if not paragraph_buffer:
            return
        p = doc.add_paragraph()
        add_inline(p, " ".join(part.strip() for part in paragraph_buffer))
        paragraph_buffer.clear()

    while index < len(lines):
        raw = lines[index]
        stripped = raw.strip()

        if not stripped:
            flush_buffer()
            index += 1
            continue
        if stripped.startswith("```"):
            flush_buffer()
            code: list[str] = []
            index += 1
            while index < len(lines) and not lines[index].strip().startswith("```"):
                code.append(lines[index])
                index += 1
            add_code_block(doc, code)
            index += 1
            continue
        if stripped.startswith("|") and index + 1 < len(lines) and re.match(r"^\|?\s*:?-+", lines[index + 1].strip()):
            flush_buffer()
            index = parse_table(lines, index, doc)
            continue
        heading = re.match(r"^(#{2,4})\s+(.+)$", stripped)
        if heading:
            flush_buffer()
            level = min(len(heading.group(1)) - 1, 3)
            p = doc.add_paragraph(style=f"Heading {level}")
            if heading.group(2).startswith("16."):
                p.paragraph_format.page_break_before = True
            add_inline(p, heading.group(2))
            index += 1
            continue
        if stripped.startswith("> "):
            flush_buffer()
            add_blockquote(doc, stripped[2:])
            index += 1
            continue
        if raw.startswith("   - "):
            flush_buffer()
            p = doc.add_paragraph(style="List Bullet 2")
            add_inline(p, stripped[2:])
            index += 1
            continue
        bullet = re.match(r"^-\s+(.+)$", stripped)
        if bullet:
            flush_buffer()
            parts = [bullet.group(1)]
            index += 1
            while index < len(lines):
                continuation = lines[index]
                if not continuation.strip():
                    break
                if continuation.startswith("  ") and not continuation.strip().startswith("-"):
                    parts.append(continuation.strip())
                    index += 1
                    continue
                break
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, " ".join(parts))
            continue
        numbered = re.match(r"^(\d+)\.\s+(.+)$", stripped)
        if numbered:
            flush_buffer()
            parts = [numbered.group(2)]
            index += 1
            while index < len(lines):
                continuation = lines[index]
                if not continuation.strip():
                    break
                if continuation.startswith("   ") and not continuation.strip().startswith("-"):
                    parts.append(continuation.strip())
                    index += 1
                    continue
                break
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.75)
            p.paragraph_format.first_line_indent = Cm(-0.48)
            p.paragraph_format.space_after = Pt(3)
            number = p.add_run(numbered.group(1) + ".  ")
            set_run_font(number, color=NAVY)
            add_inline(p, " ".join(parts))
            continue

        paragraph_buffer.append(stripped)
        index += 1

    flush_buffer()
    add_footer(doc)
    doc.core_properties.title = "RAG Assistant — подробная инструкция пользователя"
    doc.core_properties.subject = "Работа с проектами, общими и приватными документами"
    doc.core_properties.author = "НЭК"
    doc.core_properties.keywords = "RAG, инструкция, документы, проекты, источники"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
