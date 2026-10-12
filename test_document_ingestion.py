import re
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from docx.oxml import parse_xml

import document_ingestion
import docx_parser


class DocumentIngestionTests(unittest.TestCase):
    def test_short_docx_is_ingested_and_cached_without_fallback(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "short.docx"
            document = Document()
            document.add_paragraph("Код проекта: A17")
            document.save(source)
            with patch.object(document_ingestion, "CACHE_DIR", root / "cache"), \
                 patch.object(document_ingestion, "_read_with_docling") as fallback:
                result = document_ingestion.ingest_document(source)
                cached = document_ingestion.ingest_document(source)
            self.assertEqual(result.backend, "xml")
            self.assertEqual(result.text, "Код проекта: A17")
            self.assertEqual(cached.text, result.text)
            fallback.assert_not_called()

    def test_empty_docx_is_not_accepted(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "empty.docx"
            Document().save(source)
            with patch.object(document_ingestion, "CACHE_DIR", root / "cache"), \
                 patch.object(document_ingestion, "_read_with_docling", return_value=""), \
                 patch.object(document_ingestion, "_read_with_unstructured", return_value=""):
                with self.assertRaisesRegex(RuntimeError, "no readable content"):
                    document_ingestion.ingest_document(source)

    def test_location_and_unread_image_markers_are_not_content(self):
        self.assertFalse(document_ingestion.has_extracted_content("[SECTION: Chapter 1]\n[IMAGE_UNREAD: media/image1.png]"))
        self.assertTrue(document_ingestion.has_extracted_content("[SECTION: Chapter 1]\nA"))

    def test_normative_long_clause_keeps_own_number_and_body(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "Постановление 1179.docx"
            document = Document()
            document.add_paragraph("1. Утвердить правила.")
            document.add_paragraph("5. Цена второй категории определяется по формуле. " + "Подробное условие. " * 15)
            document.add_paragraph("Формула и расшифровка величин.")
            document.save(path)
            text = docx_parser.read_docx_text(path)
            self.assertIn("[SECTION: Пункт 5]", text)
            self.assertIn("Формула и расшифровка", text)
            self.assertNotIn("Пункт 1 > Пункт 5", text)

    def test_heading_level_depends_only_on_number_prefix(self):
        self.assertEqual(docx_parser._heading_level("", "", "2. Настоящее постановление вступает в силу с 1-го числа. Условие."), 1)
        self.assertEqual(docx_parser._heading_level("", "", "5.4.7 Работа с графиками"), 3)

    def test_unread_embedded_image_is_not_silently_removed(self):
        xml = parse_xml('<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><w:r><w:drawing><a:blip r:embed="rId5"/></w:drawing></w:r></w:p>')
        self.assertIn("IMAGE_UNREAD", docx_parser._extract_inline_text(xml, {"rId5": "[IMAGE_UNREAD: formula.gif]"}))

    def test_image_uses_local_ocr_and_accepts_short_meaningful_text(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "scheme.png"
            source.write_bytes(b"test-image")

            with patch.object(document_ingestion, "CACHE_DIR", root / "cache"), \
                 patch.object(document_ingestion, "_read_image_with_tesseract", return_value="ИВ0 = 42") as ocr:
                result = document_ingestion.ingest_document(source)

            self.assertEqual(result.backend, "tesseract-ocr")
            self.assertEqual(result.text, "ИВ0 = 42")
            ocr.assert_called_once_with(source)

    def test_docx_preserves_sections_tables_and_omml_fraction(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "calculation.docx"
            document = Document()
            document.add_heading("Расчёт тепловой энергии", level=1)
            paragraph = document.add_paragraph("Используется формула ")
            paragraph._p.append(parse_xml(
                '<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">'
                '<m:f><m:num><m:r><m:t>Q</m:t></m:r></m:num>'
                '<m:den><m:r><m:t>t</m:t></m:r></m:den></m:f>'
                '</m:oMath>'
            ))
            document.add_paragraph(
                "Формула применяется для расчёта удельной тепловой мощности "
                "по измеренному количеству энергии и продолжительности периода."
            )
            table = document.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "Параметр"
            table.cell(0, 1).text = "Единица"
            table.cell(1, 0).text = "Q"
            table.cell(1, 1).text = "ГДж"
            document.save(source)

            mathml_dir = root / "mathml"
            cache_dir = root / "ingestion-cache"
            with patch.object(docx_parser, "MATHML_CACHE_DIR", mathml_dir), \
                 patch.object(document_ingestion, "CACHE_DIR", cache_dir):
                extracted = docx_parser.read_docx_text(source)
                ingested = document_ingestion.ingest_document(source)
                cached = document_ingestion.ingest_document(source)

            self.assertIn("# Расчёт тепловой энергии", extracted)
            self.assertIn("[SECTION: Расчёт тепловой энергии]", extracted)
            self.assertIn("Параметр | Единица", extracted)
            self.assertIn("[MATH:", extracted)
            reference = re.search(r"\[\[MATHML_REF:([a-f0-9]{24})\]\]", extracted)
            self.assertIsNotNone(reference)
            mathml = (mathml_dir / f"{reference.group(1)}.mathml").read_text(encoding="utf-8")
            self.assertIn("<mfrac>", mathml)
            self.assertEqual(ingested.backend, "xml")
            self.assertEqual(cached.text, ingested.text)
            self.assertEqual(len(list(cache_dir.glob("*.json"))), 1)

    def test_markdown_uses_text_fallback_when_optional_parsers_fail(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "manual.md"
            source.write_text("# Руководство\n\n" + "Полезное содержимое документа. " * 8, encoding="utf-8")

            with patch.object(document_ingestion, "CACHE_DIR", root / "cache"), \
                 patch.object(document_ingestion, "_read_with_docling", side_effect=ImportError("not installed")), \
                 patch.object(document_ingestion, "_read_with_unstructured", side_effect=ImportError("not installed")):
                result = document_ingestion.ingest_document(source)

            self.assertEqual(result.backend, "xml")
            self.assertIn("# Руководство", result.text)
            self.assertEqual(len(result.warnings), 2)


if __name__ == "__main__":
    unittest.main()
