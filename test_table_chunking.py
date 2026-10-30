import unittest
import tempfile
from pathlib import Path
from unittest.mock import Mock, patch

from llama_index.core import Document
from llama_index.core.node_parser import SentenceSplitter
from llama_index.core.schema import NodeRelationship

from table_chunking import chunk_documents_with_tables
from table_context import TableContextError, parse_table_unit, protect_table_regions
from test_table_context import sample_table


class TableChunkingTests(unittest.TestCase):
    def documents(self, table, *, identity='source-version-A'):
        protected, slots = protect_table_regions(table['text'], [table])
        doc = Document(text='[SECTION: Generic]\nBefore.\n\n' + protected + '\nAfter.',
                       doc_id=identity, metadata={'file_name': 'same.docx', 'section_path': 'Generic',
                                                 'document_version': identity})
        return [doc], slots

    def split(self, documents, slots):
        return chunk_documents_with_tables(documents, slots, SentenceSplitter(chunk_size=1024, chunk_overlap=0),
                                           source_identity='physical-source-sha-A')

    def test_real_llama_splitter_preserves_all_table_units_and_source_relation(self):
        docs, slots = self.documents(sample_table())
        nodes = self.split(docs, slots)
        units = [parse_table_unit(n.text) for n in nodes if parse_table_unit(n.text)]
        self.assertEqual({x['row'] for x in units}, {0, 1})
        self.assertTrue(any('Before.' in n.text for n in nodes))
        self.assertTrue(any('After.' in n.text for n in nodes))
        for node in nodes:
            self.assertEqual(node.metadata['document_version'], 'source-version-A')
            self.assertEqual(node.relationships[NodeRelationship.SOURCE].node_id, docs[0].id_)
            self.assertNotIn('[RAG_TABLE_SLOT_', node.text)

    def test_same_filename_different_source_section_never_rebinds_headers(self):
        a, a_slots = self.documents(sample_table())
        b, b_slots = self.documents(sample_table(), identity='source-version-B')
        aa = [parse_table_unit(x.text)['table_id'] for x in self.split(a, a_slots) if parse_table_unit(x.text)]
        bb = [parse_table_unit(x.text)['table_id'] for x in self.split(b, b_slots) if parse_table_unit(x.text)]
        self.assertTrue(set(aa).isdisjoint(bb))

    def test_no_table_path_remains_exact_splitter_behavior(self):
        docs = [Document(text='Ordinary source. ' * 50, doc_id='plain')]
        splitter = SentenceSplitter(chunk_size=1024, chunk_overlap=0)
        expected = splitter.get_nodes_from_documents(docs)
        actual = chunk_documents_with_tables(docs, {}, splitter, source_identity='plain')
        self.assertEqual([x.text for x in actual], [x.text for x in expected])

    def test_lost_placeholder_is_an_error_not_silent_missing_table(self):
        docs, slots = self.documents(sample_table())
        docs[0] = Document(text='The table was accidentally removed.', doc_id=docs[0].id_, metadata=docs[0].metadata)
        with self.assertRaises(TableContextError):
            self.split(docs, slots)

    def test_duplicate_placeholder_rejected(self):
        docs, slots = self.documents(sample_table())
        docs[0] = Document(text=docs[0].text + '\n' + next(iter(slots)), doc_id=docs[0].id_, metadata=docs[0].metadata)
        with self.assertRaises(TableContextError):
            self.split(docs, slots)

    def test_long_cells_stay_structural_after_actual_token_based_splitting(self):
        table = sample_table()
        table['rows'][1]['cells'][1]['text'] = 'Длинный источник и условие. ' * 700
        docs, slots = self.documents(table)
        nodes = self.split(docs, slots)
        payloads = [parse_table_unit(n.text) for n in nodes if parse_table_unit(n.text)]
        self.assertGreater(len(payloads), 3)
        self.assertTrue(all(len(n.text) <= 2400 for n in nodes if parse_table_unit(n.text)))
        for payload in payloads:
            if payload['row'] == 1:
                self.assertTrue(any(h['text'] == 'Data source' for h in payload['headers']))


class RealIngestionTableTests(unittest.TestCase):
    def agent(self, directory):
        from llama_index.core.storage.docstore import SimpleDocumentStore
        from rag_agent import RAGAgent
        agent = object.__new__(RAGAgent)
        agent.docstore = SimpleDocumentStore()
        agent.docstore_path = Path(directory) / 'docstore.json'
        agent.vector_store = Mock()
        agent.index = Mock()
        agent._all_nodes = []
        agent.clear_cache = Mock()
        return agent

    def test_actual_docx_parser_to_ingest_to_index_preserves_source_cells(self):
        from docx import Document as WordDocument
        from docx.oxml import OxmlElement
        import document_ingestion
        with tempfile.TemporaryDirectory() as folder:
            source = Path(folder) / 'generic.docx'
            word = WordDocument()
            word.add_heading('Arbitrary table section', level=1)
            table = word.add_table(rows=2, cols=3)
            for column, value in enumerate(['Symbol', 'Data source', 'Storage']):
                table.cell(0, column).text = value
            table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
            for column, value in enumerate(['X', 'Organization Z', 'Manual entry']):
                table.cell(1, column).text = value
            word.save(source)
            agent = self.agent(folder)
            with patch.object(document_ingestion, 'CACHE_DIR', Path(folder) / 'cache'), \
                    patch('rag_agent.StorageContext.from_defaults'):
                agent.ingest(str(source))
            nodes = agent.index.insert_nodes.call_args.args[0]
            payloads = [parse_table_unit(n.text) for n in nodes if parse_table_unit(n.text)]
            body = next(x for x in payloads if x['row'] == 1)
            self.assertEqual(body['header_status'], 'explicit')
            self.assertEqual([x['text'] for x in body['cells']], ['X', 'Organization Z', 'Manual entry'])
            self.assertEqual([x['text'] for x in body['headers']], ['Symbol', 'Data source', 'Storage'])
            self.assertEqual(agent.last_ingestion_report['structured_tables'], 1)
            agent.clear_cache.assert_called_once()

    def test_actual_markdown_ingest_uses_explicit_delimiter_not_a_guessed_header(self):
        import document_ingestion
        with tempfile.TemporaryDirectory() as folder:
            source = Path(folder) / 'generic.md'
            source.write_text('# Generic\n\n| Key | Value |\n| --- | --- |\n| X | Y |', encoding='utf-8')
            agent = self.agent(folder)
            with patch.object(document_ingestion, 'CACHE_DIR', Path(folder) / 'cache'), \
                    patch('document_ingestion._read_with_docling', side_effect=ImportError), \
                    patch('document_ingestion._read_with_unstructured', side_effect=ImportError), \
                    patch('rag_agent.StorageContext.from_defaults'):
                agent.ingest(str(source))
            nodes = agent.index.insert_nodes.call_args.args[0]
            self.assertTrue(any(parse_table_unit(n.text) for n in nodes))
            self.assertEqual(agent.last_ingestion_report['structured_tables'], 1)

    def test_changed_physical_source_is_rejected_before_any_index_write(self):
        from document_ingestion import IngestedDocument
        with tempfile.TemporaryDirectory() as folder:
            source = Path(folder) / 'changing.md'
            source.write_text('Original', encoding='utf-8')
            agent = self.agent(folder)
            def read_then_change(_):
                source.write_text('Changed', encoding='utf-8')
                return IngestedDocument('Original', 'xml', [])
            with patch('rag_agent.ingest_document', side_effect=read_then_change), \
                    patch('rag_agent.StorageContext.from_defaults'):
                with self.assertRaisesRegex(ValueError, 'Source changed'):
                    agent.ingest(str(source))
            agent.index.insert_nodes.assert_not_called()
