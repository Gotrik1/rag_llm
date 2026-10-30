"""Actual DOCX -> structural rows must not change formula text or authority.

This is a transport-boundary contract, not a model-quality measurement. All
checks are deterministic and use synthetic documents without business names.
"""

import re
from unittest.mock import patch

import pytest
from docx import Document
from docx.oxml import parse_xml

import docx_parser
from grounding_verifier import build_numeric_rules, native_formula_catalogue, normalize_formula, verify_grounding
from rag_agent import ensure_requested_formula_coverage, extract_formula_lines_from_texts
from table_context import build_table_units, parse_table_unit, serialize_table_unit


QUESTION = 'Приведи формулу.'


def _evidence(text):
    return {'evidence_id': 'E1', 'source': 'synthetic.docx', 'section_path': 'Expression', 'text': text}


def _checks(answer, text):
    return [(normalize_formula(check.text), check.status) for check in
            verify_grounding(answer + ' [E1]', [_evidence(text)]).checks if check.kind == 'formula']


def _table(tmp_path, formula, *, omml=False, prefix='', row_label='R', header='Expression', max_chars=10000):
    document = Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = 'Symbol'
    table.cell(0, 1).text = header
    table.rows[0]._tr.get_or_add_trPr().append(parse_xml(
        '<w:tblHeader xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"/>'
    ))
    table.cell(1, 0).text = row_label
    if omml:
        paragraph = table.cell(1, 1).paragraphs[0]
        paragraph.add_run(prefix)
        paragraph._p.append(parse_xml(
            '<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">'
            '<m:r><m:t>R=</m:t></m:r><m:f><m:num><m:r><m:t>A</m:t></m:r></m:num>'
            '<m:den><m:r><m:t>B</m:t></m:r></m:den></m:f></m:oMath>'
        ))
    else:
        table.cell(1, 1).text = prefix + formula
    source = tmp_path / 'synthetic.docx'
    document.save(source)
    with patch.object(docx_parser, 'MATHML_CACHE_DIR', tmp_path / 'mathml'):
        parsed = docx_parser.read_docx_structured(source)
    record = parsed.tables[0]
    units = build_table_units(record, source_identity='synthetic-source-v1', tokenizer=str.split,
                              max_tokens=10000, max_chars=max_chars)
    body = next(unit for unit in units if parse_table_unit(unit)['row_kind'] == 'body')
    return parsed, record, units, body


def test_native_word_fraction_and_mathml_reference_survive_row_transport(tmp_path):
    parsed, _, _, body = _table(tmp_path, '', omml=True)
    native = native_formula_catalogue([_evidence(parsed.text)])
    structural = native_formula_catalogue([_evidence(body)])
    assert len(native) == 1
    assert structural == native
    assert re.findall(r'\[\[MATHML_REF:[a-f0-9]+\]\]', body) == re.findall(r'\[\[MATHML_REF:[a-f0-9]+\]\]', parsed.text)
    assert _checks(native[0]['display_latex'], body) == _checks(native[0]['display_latex'], parsed.text)


@pytest.mark.parametrize('formula', [
    r'$R=\frac{A}{B}$',
    r'$R=A\cdot B$',
    r'$R=\mathrm{A}_{i}+B$',
    r'$R=\frac{A_{i}}{B_{j}}$',
])
def test_latex_extraction_is_identical_after_json_row_encoding(tmp_path, formula):
    parsed, _, _, body = _table(tmp_path, formula)
    baseline = extract_formula_lines_from_texts([parsed.text], QUESTION)
    assert baseline, 'Fixture must exercise a formula accepted by the current plaintext path'
    assert extract_formula_lines_from_texts([body], QUESTION) == baseline


@pytest.mark.parametrize('formula', [r'R=\frac{A}{B}', r'R=A\cdot B', r'R=\mathrm{A}_{i}+B'])
def test_formula_matching_does_not_treat_json_escape_as_math(tmp_path, formula):
    parsed, _, _, body = _table(tmp_path, '$' + formula + '$')
    baseline = _checks('$' + formula + '$', parsed.text)
    assert baseline == [(normalize_formula(formula), 'source_matched')]
    assert _checks('$' + formula + '$', body) == baseline


def test_arbitrary_literal_word_math_with_latex_escape_keeps_source_bytes(tmp_path):
    formula = r'[MATH: R=A\cdot B]'
    parsed, _, _, body = _table(tmp_path, formula)
    baseline = extract_formula_lines_from_texts([parsed.text], QUESTION)
    assert baseline == [r'R=A\cdot B']
    assert extract_formula_lines_from_texts([body], QUESTION) == baseline
    assert _checks(r'$R=A\cdot B$', body) == _checks(r'$R=A\cdot B$', parsed.text)


def test_literal_newline_in_math_marker_is_not_backslash_n_math(tmp_path):
    formula = '[MATH: R=A+\nB]'
    parsed, _, _, body = _table(tmp_path, formula)
    baseline = _checks('$R=A+B$', parsed.text)
    assert baseline == [('R=A+B', 'source_matched')]
    assert _checks('$R=A+B$', body) == baseline


@pytest.mark.parametrize('changed', [r'R=\frac{A}{C}', r'R=\frac{A}{B}+Z', r'Q=\frac{A}{B}'])
def test_row_transport_never_authorizes_changed_terms_or_output(tmp_path, changed):
    parsed, _, _, body = _table(tmp_path, r'$R=\frac{A}{B}$')
    baseline = _checks('$' + changed + '$', parsed.text)
    assert baseline == [(normalize_formula(changed), 'unsupported')]
    assert _checks('$' + changed + '$', body) == baseline


def test_ocr_taint_cannot_be_removed_by_json_row_transport(tmp_path):
    formula = '[IMAGE_OCR: scan.png]\n[MATH: R=A+B]\n[/IMAGE_OCR]'
    parsed, _, _, body = _table(tmp_path, formula)
    baseline = _checks('$R=A+B$', parsed.text)
    assert baseline == [('R=A+B', 'unverified_ocr')]
    assert _checks('$R=A+B$', body) == baseline


def test_latex_in_verified_source_header_is_decoded_like_other_cells(tmp_path):
    parsed, _, _, body = _table(tmp_path, 'Literal data', header=r'$R=\frac{A}{B}$')
    baseline = _checks(r'$R=\frac{A}{B}$', parsed.text)
    assert baseline == [('R=(A)/(B)', 'source_matched')]
    assert _checks(r'$R=\frac{A}{B}$', body) == baseline


def test_latex_in_repeated_short_row_context_keeps_literal_backslashes(tmp_path):
    parsed, _, units, body = _table(tmp_path, r'$R=A\cdot B$', row_label='Long literal context. ' * 100,
                                    max_chars=1100)
    assert len(units) > 2
    payload = parse_table_unit(body)
    assert payload['context'] and any(r'\cdot' in cell['text'] for cell in payload['context'])
    baseline = _checks(r'$R=A\cdot B$', parsed.text)
    assert baseline == [('R=A*B', 'source_matched')]
    assert _checks(r'$R=A\cdot B$', body) == baseline


def test_disjoint_cells_are_not_stitched_into_one_native_equation(tmp_path):
    _, _, _, body = _table(tmp_path, 'B]', row_label='[MATH: R=A+')
    checks = _checks('$R=A+B$', body)
    assert all(status != 'source_matched' for _, status in checks)
    assert native_formula_catalogue([_evidence(body)]) == []


def test_header_and_data_cells_cannot_complete_each_others_equation(tmp_path):
    _, _, _, body = _table(tmp_path, 'B]', header='[MATH: R=A+')
    checks = _checks('$R=A+B$', body)
    assert all(status != 'source_matched' for _, status in checks)
    assert native_formula_catalogue([_evidence(body)]) == []


def test_invalid_unit_checksum_cannot_supply_deterministic_math_proof(tmp_path):
    _, _, _, body = _table(tmp_path, '[MATH: R=A+B]')
    invalid = body.replace('"sha256":"', '"sha256":"0', 1)
    assert parse_table_unit(invalid) is None
    assert extract_formula_lines_from_texts([invalid], QUESTION) == []
    assert native_formula_catalogue([_evidence(invalid)]) == []
    assert all(status != 'source_matched' for _, status in _checks('$R=A+B$', invalid))


@pytest.mark.parametrize('placement', ['cells', 'context'])
def test_complete_math_marker_inside_partial_cell_cannot_attest_prefix(tmp_path, placement):
    parsed, _, _, body = _table(tmp_path, '[MATH: R=A+B] [MATH: +Z]')
    assert _checks('$R=A+B$', parsed.text) == [('R=A+B', 'unsupported')]
    payload = parse_table_unit(body)
    cell = next(cell for cell in payload['cells'] if '[MATH:' in cell['text'])
    fragment = {**cell, 'text': '[MATH: R=A+B]', 'end': len('[MATH: R=A+B]')}
    if placement == 'cells':
        payload['cells'] = [fragment]
        payload['context'] = []
    else:
        payload['cells'] = [cell for cell in payload['cells'] if cell['column'] != fragment['column']]
        payload['context'] = [fragment]
    encoded = serialize_table_unit(payload)
    assert native_formula_catalogue([_evidence(encoded)]) == []
    assert all(status != 'source_matched' for _, status in _checks('$R=A+B$', encoded))


def test_numeric_rules_use_decoded_formula_with_exact_input_symbols(tmp_path):
    parsed, _, _, body = _table(tmp_path, r'$R=\frac{A}{B}$')
    baseline, diagnostics = build_numeric_rules('A=12; B=4', [_evidence(parsed.text)])
    assert len(baseline) == 1 and baseline[0]['expected_numeric_result'] == '3'
    structural, structural_diagnostics = build_numeric_rules('A=12; B=4', [_evidence(body)])
    assert structural == baseline
    assert structural_diagnostics == diagnostics


def test_structural_evidence_does_not_change_literal_answer_span_offsets(tmp_path):
    _, _, _, body = _table(tmp_path, r'$R=\frac{A}{B}$')
    answer = 'An introduction.\n\n$$R=\\frac{A}{B}$$ [E1]'
    checks = [check for check in verify_grounding(answer, [_evidence(body)]).checks if check.kind == 'formula']
    assert len(checks) == 1 and checks[0].status == 'source_matched'
    check = checks[0]
    assert answer[check.start:check.end] == r'$$R=\frac{A}{B}$$'


@pytest.mark.parametrize('formula', [r'R=\frac{A}{B}', r'\frac{A}{B}\cdot C'])
def test_automatic_formula_completion_confirms_decoded_source_not_wire_escaping(tmp_path, formula):
    parsed, _, _, body = _table(tmp_path, formula)
    answer = 'A literal explanation.'
    baseline = ensure_requested_formula_coverage(answer, QUESTION, [formula], [_evidence(parsed.text)])
    assert baseline != answer
    assert ensure_requested_formula_coverage(answer, QUESTION, [formula], [_evidence(body)]) == baseline
